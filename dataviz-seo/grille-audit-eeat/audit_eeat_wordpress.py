#!/usr/bin/env python3
"""
Audit E-E-A-T en lot pour un blog WordPress.

Récupère les articles publiés via l'API REST de WordPress (ou un fichier JSON),
applique les contrôles automatiques de la grille pondérée, puis exporte un
rapport Excel pré-rempli : notes suggérées, score, verdict et signaux détectés.
La note finale reste humaine : relisez et corrigez les cellules jaunes.

Usage :
    python audit_eeat_wordpress.py https://monblog.fr
    python audit_eeat_wordpress.py https://monblog.fr --max 50 --out rapport.xlsx
    python audit_eeat_wordpress.py --json articles.json          # hors ligne

Dépendances : Python 3.9+, openpyxl  (pip install openpyxl)
"""
import argparse
import html
import json
import re
import sys
import time
import urllib.error
import urllib.request
from datetime import date, datetime
from html.parser import HTMLParser

# --------------------------------------------------------------------------
# Paramètres de la grille (modifiables aussi dans l'onglet « Paramètres »)
# --------------------------------------------------------------------------
CRITERIA = [
    ("exp", "Preuve d'expérience", 25),
    ("src", "Traçabilité des sources", 20),
    ("time", "Cohérence temporelle", 15),
    ("struct", "Intégrité de la structure", 15),
    ("dens", "Densité et valeur ajoutée", 15),
    ("auth", "Signaux auteur et site", 10),
]
SEUIL_FIABLE, SEUIL_VERIFIER, PLAFOND = 75, 50, 40

PLUGINS = ["yoast", "rank math", "seopress", "all in one seo", "aioseo", "the seo framework",
           "slim seo", "squirrly", "wp rocket", "litespeed cache", "w3 total cache", "wp super cache",
           "flyingpress", "perfmatters", "autoptimize", "asset cleanup", "wp fastest cache", "nitropack",
           "elementor", "jetpack", "wordfence", "schema pro", "smush", "shortpixel", "imagify", "ewww",
           "redirection", "woocommerce", "gutenberg", "kadence", "generatepress", "astra",
           "query monitor", "site kit", "polylang", "wpml", "updraftplus", "cloudflare",
           "search console", "pagespeed", "lighthouse", "screaming frog", "semrush", "ahrefs"]
BRANDS = ["semrush", "ahrefs", "sparktoro", "backlinko", "moz", "hubspot", "similarweb", "brightedge",
          "gartner", "statista", "sistrix", "seoclarity", "authoritas", "datos", "google", "bing",
          "microsoft", "openai", "perplexity"]
HEDGES = ["varie selon les études", "ordre de grandeur", "n'est pas détaillée", "ne sont pas détaillés",
          "selon certaines études", "selon plusieurs études", "selon des experts", "on estime que",
          "les études montrent", "certaines sources", "il semblerait", "selon les sources",
          "à titre indicatif"]
CLICHES = ["dans le paysage actuel", "il est crucial", "il est essentiel", "en constante évolution",
           "à l'ère de", "levier incontournable", "game changer", "n'est plus une option", "plus que jamais",
           "en somme", "en définitive", "pierre angulaire", "au cœur de", "dans un monde où",
           "il convient de", "force est de constater", "optimiser votre visibilité",
           "transforme la stratégie en acte", "véritable atout", "booster votre", "boostez",
           "incontournable", "révolutionne", "sans plus attendre", "le saviez-vous", "en conclusion",
           "il est important de noter", "jouer un rôle clé", "clé du succès"]
MONTHS = ["janvier", "février", "mars", "avril", "mai", "juin", "juillet", "août", "septembre",
          "octobre", "novembre", "décembre"]

RE_FIRST = re.compile(
    r"\b(j'ai (testé|utilisé|installé|essayé|renoncé|désactivé|activé|mesuré|constaté|migré|remplacé|"
    r"configuré|abandonné)|nous avons (testé|utilisé|installé|mesuré|constaté|migré)|dans mon expérience|"
    r"chez mes clients|sur mes sites|sur un de mes sites|mon client|j'utilise|je recommande|je n'utilise plus)",
    re.I)
RE_STAT = re.compile(r"\d+[,.]?\d*\s?%|\b\d+\s?(fois|millions?|milliards?)\b|\bx\d+\b", re.I)
RE_METRIC = re.compile(r"\d+[,.]?\d*\s?(s|ms|%|secondes?|clics|visites|positions?)\b|passé de|avant.*après", re.I)
RE_VERSION = re.compile(r"\b\d+\.\d+(\.\d+)?\b")
RE_STUDY = re.compile(r"\b(étude|données|rapport|baromètre|enquête|tendances?|statistiques?|analyse|bilan|chiffres)\b", re.I)
RE_FULLYEAR = re.compile(r"(sur l'ensemble (des données|de l'année)|consolidé|bilan (annuel|de l'année)|"
                         r"sur toute l'année|année complète|sur l'année)", re.I)
RE_WORD = re.compile(r"[\w'-]+", re.U)


# --------------------------------------------------------------------------
# Récupération des articles
# --------------------------------------------------------------------------
def fetch_posts(site, max_posts=50, pause=0.5):
    """Récupère les articles publiés via /wp-json/wp/v2/posts (pagination incluse)."""
    base = site.rstrip("/") + "/wp-json/wp/v2/posts"
    posts, page = [], 1
    while len(posts) < max_posts:
        per_page = min(100, max_posts - len(posts))
        url = f"{base}?per_page={per_page}&page={page}&status=publish&_embed=author"
        req = urllib.request.Request(url, headers={"User-Agent": "audit-eeat/1.0", "Accept": "application/json"})
        try:
            with urllib.request.urlopen(req, timeout=30) as r:
                batch = json.loads(r.read().decode("utf-8"))
                total_pages = int(r.headers.get("X-WP-TotalPages", page))
        except urllib.error.HTTPError as e:
            if e.code == 400 and page > 1:
                break  # page au-delà de la dernière
            sys.exit(f"Erreur HTTP {e.code} sur {url}. L'API REST est-elle désactivée ou protégée ?")
        except urllib.error.URLError as e:
            sys.exit(f"Impossible de joindre {site} : {e.reason}")
        if not batch:
            break
        posts.extend(batch)
        print(f"  page {page}/{total_pages} : {len(batch)} article(s)")
        if page >= total_pages:
            break
        page += 1
        time.sleep(pause)  # politesse envers le serveur
    return posts[:max_posts]


def normalize_post(p):
    """Ramène un article de l'API (ou d'un JSON simplifié) à un format commun."""
    def val(x):
        return x.get("rendered", "") if isinstance(x, dict) else (x or "")
    author = ""
    emb = p.get("_embedded", {}).get("author") or []
    if emb and isinstance(emb[0], dict):
        author = emb[0].get("name", "")
    author = author or p.get("author_name", "") or ""
    return {
        "url": p.get("link", ""),
        "title": html.unescape(re.sub(r"<[^>]+>", "", val(p.get("title")))).strip(),
        "date": (p.get("date") or "")[:10],
        "modified": (p.get("modified") or "")[:10],
        "author": author,
        "html": val(p.get("content")),
    }


# --------------------------------------------------------------------------
# Découpage du HTML en blocs (titres / paragraphes / liens)
# --------------------------------------------------------------------------
class BlockParser(HTMLParser):
    BLOCKS = {"h1", "h2", "h3", "h4", "h5", "h6", "p", "li", "td", "blockquote", "figcaption"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.blocks, self.stack, self.cur, self.in_a, self.a_text = [], [], None, False, ""

    def handle_starttag(self, tag, attrs):
        if tag in self.BLOCKS:
            if self.cur is not None:
                self.stack.append(self.cur)
            self.cur = {"type": "h" if tag[0] == "h" and tag[1:].isdigit() else "p", "text": "", "links": []}
        elif tag == "a" and dict(attrs).get("href"):
            self.in_a, self.a_text = True, ""
        elif tag == "br" and self.cur is not None:
            self.cur["text"] += " "

    def handle_endtag(self, tag):
        if tag == "a" and self.in_a:
            self.in_a = False
            if self.cur is not None and self.a_text.strip():
                self.cur["links"].append(self.a_text.strip())
        elif tag in self.BLOCKS and self.cur is not None:
            self.cur["text"] = re.sub(r"\s+", " ", self.cur["text"]).strip()
            if self.cur["text"]:
                self.blocks.append(self.cur)
            self.cur = self.stack.pop() if self.stack else None

    def handle_data(self, data):
        if self.cur is not None:
            self.cur["text"] += data
        if self.in_a:
            self.a_text += data


def norm(s):
    return s.replace("’", "'").replace("‘", "'").replace(" ", " ")


def parse_blocks(content):
    content = norm(content)
    if re.search(r"</?(p|h[1-6]|li|div)[\s>/]", content, re.I):
        bp = BlockParser()
        bp.feed(content)
        return bp.blocks
    # Texte brut : on devine les titres (lignes courtes sans ponctuation finale)
    lines, blocks = content.split("\n"), []
    for i, l in enumerate(lines):
        t = l.strip()
        if not t:
            continue
        nxt = next((x for x in lines[i + 1:] if x.strip()), "")
        heading = t.startswith("#") or (len(RE_WORD.findall(t)) <= 12 and not re.search(r"[.!?:;,…]$", t)
                                        and not re.match(r"^[-*•\d]", t) and nxt)
        blocks.append({"type": "h" if heading else "p", "text": t.lstrip("# "), "links": re.findall(r"https?://\S+", t)})
    return blocks


def sentences(text):
    return [s.strip() for s in re.split(r"(?<=[.!?…])\s+|\n+", text) if len(s.strip()) > 2]


def words(s):
    return RE_WORD.findall(s)


# --------------------------------------------------------------------------
# Contrôles automatiques (mêmes règles que la page interactive)
# --------------------------------------------------------------------------
def analyze(post):
    blocks = parse_blocks(post["html"])
    sents = [(s, b) for b in blocks if b["type"] == "p" for s in sentences(b["text"])]
    body = "\n".join(b["text"] for b in blocks)
    n_words = len(words(body))
    try:
        pub = datetime.strptime(post["date"], "%Y-%m-%d").date()
    except ValueError:
        pub = date.today()
    F = {k: [] for k, _, _ in CRITERIA}  # signaux par critère : (niveau, signal, extrait)
    S, why = {}, {}

    # 1. Expérience
    anon = named = metric = 0
    for i, (s, _) in enumerate(sents):
        if not RE_FIRST.search(s):
            continue
        ctx = (s + " " + (sents[i + 1][0] if i + 1 < len(sents) else "")).lower()
        has_name = any(p in ctx for p in PLUGINS) or RE_VERSION.search(ctx)
        has_metric = RE_METRIC.search(ctx)
        if has_name:
            named += 1
            metric += bool(has_metric)
            F["exp"].append(("ok" if has_metric else "moyen",
                             "Expérience nommée et chiffrée" if has_metric else "Outil nommé, sans mesure", s))
        else:
            anon += 1
            F["exp"].append(("alerte", "Première personne, outil jamais nommé", s))
    if not anon and not named:
        S["exp"], why["exp"] = 1, "Aucun témoignage à la première personne : chercher captures ou cas clients."
    elif anon and not named:
        S["exp"], why["exp"] = 0, f"{anon} témoignage(s) sans nom d'outil ni de site."
    elif metric and not anon:
        S["exp"], why["exp"] = 2, "Outils nommés et résultats chiffrés."
    else:
        S["exp"], why["exp"] = 1, f"{named} expérience(s) nommée(s), {anon} anonyme(s), {metric} chiffrée(s)."

    # 2. Sources
    stats = unlinked = hedges = 0
    for s, b in sents:
        low = s.lower()
        for h in HEDGES:
            if h in low:
                hedges += 1
                F["src"].append(("alerte", "Formule d'auto-protection", s))
        if RE_STAT.search(s):
            stats += 1
            linked = any(a in s for a in b["links"]) or "http" in s
            if not linked:
                unlinked += 1
                brand = next((x for x in BRANDS if re.search(r"\b" + x + r"\b", low)), None)
                F["src"].append(("moyen" if brand else "alerte",
                                 f"Chiffre attribué à {brand.capitalize()} sans lien" if brand else "Chiffre sans source", s))
    if not stats:
        S["src"], why["src"] = 1, "Aucun chiffre détecté : vérifier les affirmations à la main."
    elif not unlinked and not hedges:
        S["src"], why["src"] = 2, f"{stats} chiffre(s), tous liés."
    elif unlinked / stats > 0.5 or hedges >= 2:
        S["src"], why["src"] = 0, f"{unlinked}/{stats} chiffre(s) sans lien, {hedges} formule(s) d'auto-protection."
    else:
        S["src"], why["src"] = 1, f"{unlinked}/{stats} chiffre(s) sans lien."

    # 3. Temporel
    impossible = risky = 0
    for s, _ in sents:
        low = s.lower()
        for y in map(int, re.findall(r"\b((?:19|20)\d{2})\b", s)):
            if y > pub.year:
                impossible += 1
                F["time"].append(("alerte", f"Année {y} postérieure à la publication", s))
            elif y == pub.year:
                m = next((i for i, mo in enumerate(MONTHS) if re.search(mo + r"\s+" + str(y), low)), None)
                if m is not None and m + 1 > pub.month:
                    impossible += 1
                    F["time"].append(("alerte", f"{MONTHS[m]} {y} est après la publication", s))
                elif RE_FULLYEAR.search(s):
                    impossible += 1
                    F["time"].append(("alerte", f"Données « annuelles » {y} alors que l'année n'est pas finie", s))
                elif RE_STUDY.search(s):
                    risky += 1
                    F["time"].append(("moyen", f"Étude de l'année en cours ({y}) à vérifier", s))
    S["time"] = 0 if impossible else (1 if risky else 2)
    why["time"] = (f"{impossible} date(s) impossible(s)." if impossible else
                   f"{risky} étude(s) de {pub.year} à vérifier." if risky else "Aucune date incohérente.")

    # 4. Structure
    heads = []
    for i, b in enumerate(blocks):
        if b["type"] == "h":
            w = 0
            for nb in blocks[i + 1:]:
                if nb["type"] == "h":
                    break
                w += len(words(nb["text"]))
            heads.append((b["text"], w))
    thin = 0
    for t, w in heads:
        if w < 50 and not re.search(r"conclusion", t, re.I):
            thin += 1
            F["struct"].append(("alerte", f"Section de {w} mots", t))
        m = re.search(r"\b(\d+)\s+(outils|extensions|plugins|étapes|astuces|conseils|erreurs)\b", t, re.I)
        if m:
            F["struct"].append(("moyen", f"Promet {m.group(1)} {m.group(2)} : compter les éléments", t))
        if re.search(r"\b(liste|top|comparatif)\b", t, re.I) and w < 80:
            F["struct"].append(("alerte", "Liste annoncée quasi vide", t))
    if body.strip() and not re.search(r"[.!?…»)\]]$", body.strip()):
        F["struct"].append(("alerte", "Fin de texte tronquée", body.strip()[-120:]))
    bad = sum(1 for f in F["struct"] if f[0] == "alerte")
    S["struct"] = 1 if not heads else (0 if bad >= 2 else (1 if F["struct"] else 2))
    why["struct"] = f"{len(heads)} titre(s), {thin} section(s) de moins de 50 mots." if heads else "Aucun titre détecté."

    # 5. Densité
    cl = 0
    for s, _ in sents:
        hit = next((c for c in CLICHES if c in s.lower()), None)
        if hit:
            cl += 1
            F["dens"].append(("moyen", f"Formule creuse : {hit}", s))
    long_s = [s for s, _ in sents if len(words(s)) >= 8]
    sets = [{w for w in words(s.lower()) if len(w) > 3} for s in long_s]
    rep = 0
    for i in range(len(sets)):
        for j in range(i + 1, len(sets)):
            inter = len(sets[i] & sets[j])
            union = len(sets[i] | sets[j]) or 1
            if inter / union >= 0.5:
                rep += 1
                F["dens"].append(("alerte", f"Phrase répétée ({round(inter / union * 100)} %)", long_s[j]))
    per1k = (cl + rep * 2) / n_words * 1000 if n_words else 0
    S["dens"] = 1 if not n_words else (0 if per1k >= 12 else (1 if per1k >= 4 else 2))
    why["dens"] = f"{cl} formule(s) creuse(s), {rep} répétition(s) pour {n_words} mots."

    # 6. Auteur : non détectable dans le texte
    S["auth"], why["auth"] = None, "À vérifier sur le site (bio, profil, mentions légales)."

    return {"scores": S, "why": why, "findings": F, "words": n_words, "impossible": impossible,
            "headings": len(heads)}


# --------------------------------------------------------------------------
# Export Excel
# --------------------------------------------------------------------------
def export_xlsx(rows, out, site):
    from openpyxl import Workbook
    from openpyxl.comments import Comment
    from openpyxl.formatting.rule import CellIsRule, FormulaRule
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter as L
    from openpyxl.worksheet.datavalidation import DataValidation

    FN, TEAL = "Arial", "0E6B63"
    f = lambda **k: Font(name=FN, **k)
    side = Side(style="thin", color="D5DDDB")
    box = Border(left=side, right=side, top=side, bottom=side)
    head = PatternFill("solid", fgColor=TEAL)
    inp = PatternFill("solid", fgColor="FFF7D6")
    red, amb, grn = (PatternFill("solid", fgColor=c) for c in ("F7E0DD", "F7ECD9", "E3F1E4"))
    ctr = Alignment(horizontal="center", vertical="top", wrap_text=True)
    wrap = Alignment(wrap_text=True, vertical="top")

    wb = Workbook()
    A = wb.active
    A.title = "Audit"

    # Paramètres
    P = wb.create_sheet("Paramètres")
    P["A1"] = "Paramètres de la grille"; P["A1"].font = f(bold=True, size=14, color=TEAL)
    P["A2"] = "Modifiez les cellules jaunes : tout le classeur se recalcule."; P["A2"].font = f(italic=True, color="586866")
    for c, h in enumerate(["Critère", "Poids"], 1):
        P.cell(4, c, h).font = f(bold=True, color="FFFFFF"); P.cell(4, c).fill = head
    for i, (_, name, w) in enumerate(CRITERIA):
        P.cell(5 + i, 1, name).font = f()
        P.cell(5 + i, 2, w).font = f(color="0000FF"); P.cell(5 + i, 2).fill = inp
    P["A11"] = "Total des poids"; P["B11"] = "=SUM(B5:B10)"
    for r, (lab, v) in enumerate([("Seuil « Fiable » (≥)", SEUIL_FIABLE), ("Seuil « À vérifier » (≥)", SEUIL_VERIFIER),
                                  ("Plafond si critère éliminatoire", PLAFOND),
                                  ("Note par défaut « Auteur / site » tant qu'elle n'est pas saisie", 1)], 13):
        P.cell(r, 1, lab).font = f(); P.cell(r, 2, v).font = f(color="0000FF"); P.cell(r, 2).fill = inp
    P.column_dimensions["A"].width = 58; P.column_dimensions["B"].width = 10

    # Audit
    A["A1"] = "Audit E-E-A-T en lot"; A["A1"].font = f(bold=True, size=16, color=TEAL)
    A["A2"] = (f"Source : {site} · extrait le {date.today():%d/%m/%Y}. Notes pré-remplies par le script : "
               "relisez et corrigez les cellules jaunes. « Auteur / site », « Stat introuvable » et "
               "« Faux témoignage » demandent une vérification humaine.")
    A["A2"].font = f(italic=True, color="586866"); A.merge_cells("A2:U2"); A["A2"].alignment = wrap
    A.row_dimensions[2].height = 32
    heads = ["N°", "URL", "Titre", "Publié le", "Modifié le", "Auteur", "Mots",
             "Expérience", "Sources", "Temporel", "Structure", "Densité", "Auteur / site",
             "Date impossible", "Stat introuvable", "Faux témoignage",
             "Score /100", "Verdict", "Statut", "Action", "Principaux signaux"]
    for c, h in enumerate(heads, 1):
        cell = A.cell(4, c, h)
        cell.font = f(bold=True, color="FFFFFF"); cell.fill = head; cell.alignment = ctr; cell.border = box
    A["H3"] = "Notes 0-2 (suggérées, à confirmer)"; A.merge_cells("H3:M3")
    A["N3"] = "Éliminatoires (Oui / Non)"; A.merge_cells("N3:P3")
    for c in ("H3", "N3"):
        A[c].font = f(bold=True, color=TEAL); A[c].alignment = Alignment(horizontal="center")

    first = 5
    last = first + max(len(rows), 1) - 1
    for i, (post, res) in enumerate(rows):
        r = first + i
        sc = res["scores"]
        top = [x for k, _, _ in CRITERIA for x in res["findings"][k] if x[0] == "alerte"][:3]
        vals = [i + 1, post["url"], post["title"], _d(post["date"]), _d(post["modified"]), post["author"],
                res["words"], sc["exp"], sc["src"], sc["time"], sc["struct"], sc["dens"], None,
                "Oui" if res["impossible"] else "Non", None, None]
        for c, v in enumerate(vals, 1):
            A.cell(r, c, v)
        brut = (f'ROUND((H{r}*Paramètres!$B$5+I{r}*Paramètres!$B$6+J{r}*Paramètres!$B$7+K{r}*Paramètres!$B$8'
                f'+L{r}*Paramètres!$B$9+IF(M{r}="",Paramètres!$B$16,M{r})*Paramètres!$B$10)'
                f'/2/Paramètres!$B$11*100,1)')
        A.cell(r, 17, f'=IF(COUNTIF(N{r}:P{r},"Oui")>0,MIN(Paramètres!$B$15,{brut}),{brut})')
        A.cell(r, 18, f'=IF(Q{r}>=Paramètres!$B$13,"Fiable",IF(Q{r}>=Paramètres!$B$14,"À vérifier","Suspect"))')
        A.cell(r, 19, f'=IF(OR(M{r}="",O{r}="",P{r}=""),"Provisoire","Validé")')
        A.cell(r, 20, f'=IF(R{r}="Fiable","Conserver",IF(R{r}="À vérifier","Enrichir : preuve d\'expérience + liens",'
                      f'"Réécrire ou dépublier"))')
        A.cell(r, 21, " | ".join(f"{k}" for _, k, _ in top) or "Aucune alerte automatique")
        for c in range(1, 22):
            cell = A.cell(r, c)
            cell.border = box; cell.font = f(size=10)
            cell.alignment = ctr if 4 <= c <= 20 else wrap
            if 8 <= c <= 16:
                cell.fill = inp
        A.cell(r, 4).number_format = A.cell(r, 5).number_format = "DD/MM/YYYY"
        A.cell(r, 17).number_format = "0.0"; A.cell(r, 17).font = f(size=10, bold=True)
        for j, (k, _, _) in enumerate(CRITERIA):
            A.cell(r, 8 + j).comment = Comment(res["why"][k], "Script") if res["why"][k] else None

    dv_n = DataValidation(type="whole", operator="between", formula1="0", formula2="2", allow_blank=True,
                          showErrorMessage=True, error="Saisissez 0, 1 ou 2.")
    dv_y = DataValidation(type="list", formula1='"Oui,Non"', allow_blank=True)
    A.add_data_validation(dv_n); A.add_data_validation(dv_y)
    dv_n.add(f"H{first}:M{last}"); dv_y.add(f"N{first}:P{last}")
    rng = f"H{first}:M{last}"
    A.conditional_formatting.add(rng, FormulaRule(formula=[f'AND(H{first}<>"",H{first}=0)'], fill=red,
                                                  font=Font(name=FN, bold=True, color="B63326")))
    A.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=["1"], fill=amb))
    A.conditional_formatting.add(rng, CellIsRule(operator="equal", formula=["2"], fill=grn))
    A.conditional_formatting.add(f"N{first}:P{last}", CellIsRule(operator="equal", formula=['"Oui"'], fill=red,
                                                                 font=Font(name=FN, bold=True, color="B63326")))
    for label, fill, col in (("Suspect", red, "B63326"), ("À vérifier", amb, "9A6414"), ("Fiable", grn, "2E7D32")):
        A.conditional_formatting.add(f"Q{first}:R{last}", FormulaRule(formula=[f'$R{first}="{label}"'], fill=fill,
                                                                      font=Font(name=FN, bold=True, color=col)))
    for c, w in enumerate([5, 34, 40, 11, 11, 14, 7, 10, 9, 9, 10, 9, 10, 10, 10, 10, 9, 11, 11, 22, 50], 1):
        A.column_dimensions[L(c)].width = w
    A.freeze_panes = "D5"
    A.auto_filter.ref = f"A4:U{last}"

    # Signaux (détail)
    D = wb.create_sheet("Signaux")
    for c, h in enumerate(["N°", "Titre", "Critère", "Niveau", "Signal", "Extrait"], 1):
        D.cell(1, c, h).font = f(bold=True, color="FFFFFF"); D.cell(1, c).fill = head
    r = 2
    for i, (post, res) in enumerate(rows):
        for k, name, _ in CRITERIA:
            for lvl, sig, txt in res["findings"][k]:
                for c, v in enumerate([i + 1, post["title"], name, lvl, sig, txt[:300]], 1):
                    D.cell(r, c, v).font = f(size=10); D.cell(r, c).alignment = wrap
                r += 1
    D.conditional_formatting.add(f"D2:D{max(r, 2)}", CellIsRule(operator="equal", formula=['"alerte"'], fill=red))
    D.conditional_formatting.add(f"D2:D{max(r, 2)}", CellIsRule(operator="equal", formula=['"moyen"'], fill=amb))
    D.conditional_formatting.add(f"D2:D{max(r, 2)}", CellIsRule(operator="equal", formula=['"ok"'], fill=grn))
    for c, w in enumerate([5, 36, 24, 9, 40, 80], 1):
        D.column_dimensions[L(c)].width = w
    D.freeze_panes = "A2"
    D.auto_filter.ref = f"A1:F{max(r - 1, 1)}"

    # Synthèse
    S = wb.create_sheet("Synthèse", 1)
    S["A1"] = "Rapport de conformité E-E-A-T"; S["A1"].font = f(bold=True, size=16, color=TEAL)
    S["A2"] = f"{site} · {date.today():%d/%m/%Y}"; S["A2"].font = f(italic=True, color="586866")
    rg = lambda col: f"Audit!{col}{first}:{col}{last}"
    items = [("Articles audités", f"=COUNT({rg('Q')})", "0"),
             ("Score moyen /100", f'=IFERROR(AVERAGE({rg("Q")}),"-")', "0.0"),
             ("Fiables", f'=COUNTIF({rg("R")},"Fiable")', "0"),
             ("À vérifier", f'=COUNTIF({rg("R")},"À vérifier")', "0"),
             ("Suspects", f'=COUNTIF({rg("R")},"Suspect")', "0"),
             ("Part de suspects", "=IFERROR(B8/B4,0)", "0.0%"),
             ("Avec au moins un critère éliminatoire",
              f'=SUMPRODUCT(--((({rg("N")}="Oui")+({rg("O")}="Oui")+({rg("P")}="Oui"))>0))', "0"),
             ("Audits encore provisoires", f'=COUNTIF({rg("S")},"Provisoire")', "0")]
    for c, h in enumerate(["Indicateur", "Valeur"], 1):
        S.cell(3, c, h).font = f(bold=True, color="FFFFFF"); S.cell(3, c).fill = head
    for i, (lab, fm, nf) in enumerate(items, 4):
        S.cell(i, 1, lab).font = f(); S.cell(i, 2, fm).number_format = nf; S.cell(i, 2).font = f()
    for c, h in enumerate(["Critère", "Moyenne /2", "Articles à 0"], 1):
        S.cell(13, c, h).font = f(bold=True, color="FFFFFF"); S.cell(13, c).fill = head
    for j, col in enumerate("HIJKLM"):
        S.cell(14 + j, 1, f"=Paramètres!A{5 + j}").font = f()
        S.cell(14 + j, 2, f'=IFERROR(AVERAGE({rg(col)}),"-")').number_format = "0.00"
        S.cell(14 + j, 3, f"=COUNTIF({rg(col)},0)")
    S.conditional_formatting.add("B14:B19", CellIsRule(operator="lessThan", formula=["1"], fill=red))
    S.column_dimensions["A"].width = 42; S.column_dimensions["B"].width = 12; S.column_dimensions["C"].width = 13

    wb.save(out)


def _d(s):
    try:
        return datetime.strptime(s, "%Y-%m-%d").date()
    except (TypeError, ValueError):
        return None


# --------------------------------------------------------------------------
def main():
    ap = argparse.ArgumentParser(description="Audit E-E-A-T en lot d'un blog WordPress.")
    ap.add_argument("site", nargs="?", help="URL du site WordPress, ex. https://monblog.fr")
    ap.add_argument("--json", help="Fichier JSON d'articles (format de l'API WordPress) au lieu du site")
    ap.add_argument("--max", type=int, default=50, help="Nombre maximum d'articles (défaut : 50)")
    ap.add_argument("--out", default=f"audit-eeat-{date.today():%Y-%m-%d}.xlsx", help="Fichier Excel de sortie")
    a = ap.parse_args()
    if not a.site and not a.json:
        ap.error("indiquez l'URL du site ou --json fichier.json")

    if a.json:
        with open(a.json, encoding="utf-8") as fh:
            raw = json.load(fh)[: a.max]
        source = a.json
    else:
        print(f"Récupération des articles de {a.site} …")
        raw = fetch_posts(a.site, a.max)
        source = a.site
    if not raw:
        sys.exit("Aucun article trouvé.")

    rows = []
    for p in raw:
        post = normalize_post(p)
        rows.append((post, analyze(post)))

    export_xlsx(rows, a.out, source)
    print(f"\n{len(rows)} article(s) audité(s) → {a.out}")
    print("Notes les plus basses (suggestions, avant relecture humaine) :")
    ranked = sorted(rows, key=lambda x: sum((x[1]["scores"][k] or 0) * w for k, _, w in CRITERIA))
    for post, res in ranked[:5]:
        flag = " [date impossible]" if res["impossible"] else ""
        print(f"  - {post['title'][:70]}{flag}")


if __name__ == "__main__":
    main()
